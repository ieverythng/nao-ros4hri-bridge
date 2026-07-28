#!/usr/bin/env python3
"""Render the model-invariance findings into the retained design-report template."""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Pt

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from build_model_agnostic_report import (  # noqa: E402
    _add_bullets,
    _replace_template_header,
    _set_cell_shading,
    _set_cell_text,
    _set_paragraph,
    _set_style_fonts,
    _table,
)


DATE = "21 July 2026"


def _paragraph(doc: Document, index: int, text: str) -> None:
    _set_paragraph(doc.paragraphs[index], text)


def _rate_text(rate: dict) -> str:
    return f"{rate.get('pass', 0)}/{rate.get('scored', 0)}"


def _replace_header_text(output: Path) -> None:
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".docx", delete=False) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(output, "r") as source, zipfile.ZipFile(temporary_path, "w") as target:
            for info in source.infolist():
                content = source.read(info.filename)
                if info.filename == "word/header1.xml":
                    content = content.replace(b"Model-Agnostic Runtime Qualification", b"Model Invariance Runtime Qualification")
                    content = content.replace(b"20 July 2026", b"21 July 2026")
                target.writestr(info, content)
        os.replace(temporary_path, output)
    finally:
        if temporary_path.exists():
            temporary_path.unlink()


def build(template: Path, summary_path: Path, output: Path) -> None:
    summary = __import__("json").loads(summary_path.read_text(encoding="utf-8"))
    models = {item["model"]: item for item in summary["models"]}
    doc = Document(str(template))
    _set_style_fonts(doc)

    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if "Report title" in paragraph.text:
                _set_paragraph(paragraph, "Model Invariance Runtime Qualification")
            elif paragraph.text.strip() == "Date":
                _set_paragraph(paragraph, DATE)

    cover = doc.tables[0]
    _set_cell_text(cover.cell(0, 0), "Frozen NAO ROS4HRI E2E model comparison", size=10)
    _set_cell_text(cover.cell(0, 2), "Prepared for TFM review\n" + DATE, size=9)
    overview = doc.tables[1]
    for index, value in enumerate(("Theme", "Observation", "Implication")):
        _set_cell_text(overview.cell(0, index), value, bold=True, size=8.5)
        _set_cell_shading(overview.cell(0, index), "D9E2F3")
    overview_rows = [
        ("Endpoint", "vLLM /v1/models unavailable", "No vLLM semantic score was fabricated"),
        ("Cells", "Gemma4 31B, Nemotron 3 Super, Gemma4 cloud", "Three Ollama models passed startup preflight"),
        ("Evidence", f"{summary['total_case_records']} retained case rows", "Raw JSON and derived thesis dataset preserved"),
    ]
    for row_index, values in enumerate(overview_rows, start=1):
        for column_index, value in enumerate(values):
            _set_cell_text(overview.cell(row_index, column_index), value, size=8.2)

    _paragraph(doc, 2, "Model Invariance Runtime Qualification")
    _paragraph(doc, 5, "Executive summary | Introduction | Key findings | Conditions | Evidence patterns | Implications | Recommendations | Conclusion | Appendix")
    _paragraph(doc, 7, "The frozen NAO ROS4HRI stack passed the required startup gate with three Ollama models. The requested vLLM endpoint was unavailable during the final inventory probe, so it is recorded as a blocked semantic cell.")
    _paragraph(doc, 8, "The two Gemma variants reproduced the same main and extreme-capability failure shape while passing the environment, KB, and deterministic fake-skill success families. Nemotron retained environment execution and targeted fail-once recovery but introduced additional planner coverage, KB wording, trace-correlation, and report-closure variance.")
    _paragraph(doc, 10, "Three model cells")
    _paragraph(doc, 11, "Gemma4 variants shared the failure intersection; this is the primary evidence against attributing every failure to model quality.")
    _paragraph(doc, 12, "No vLLM semantic result was scored because /v1/models returned HTTP 000.")
    _paragraph(doc, 14, "The comparison keeps the frozen image, ROS graph, prompt and generation settings, KnowledgeCore fixtures, fake-skill policies, questionnaire cases, and speech evidence policy constant. The only experimental factor is the selected Ollama model, with the declared no-driver profile used for the two alternative cells after the external NAOqi endpoint timed out.")
    _paragraph(doc, 15, "Live questionnaire rows are runtime stress evidence over fixed cases. They are distinct from deterministic source and harness tests. A pass proves the expected evidence path for that case, not universal success over all natural-language inputs.")
    _paragraph(doc, 17, "The strongest invariant is the 11/11 environment result for all three models. The Gemma variants also passed 7/7 KB stress and 9/9 all-success fake-deep cases. Nemotron passed 4/6 scoreable KB cases and 8/9 fake-deep cases, with its targeted fail-once navigation probe passing with observed failure and replan.")
    _paragraph(doc, 19, "The standard-family strict pass rates were 84.7% for gemma4:31b-cloud, 85.0% for gemma4:cloud, and 76.8% for nemotron-3-super:cloud. These are one-run engineering rates, not a universal model ranking.")
    _paragraph(doc, 21, "Shared failures concentrated in generic pick admission, maximal kitchen delivery, missing-recipient clarification, universal target selection, and high-composition capability coverage. Nemotron-only variance concentrated in planner coverage, KB speech projection, trace correlation, and terminal report closure.")
    _paragraph(doc, 23, "Key takeaway. The stack contract remained usable across the tested models, while generated route, target, plan, wording, and closure behavior remained model-dependent.")
    _paragraph(doc, 25, "Implications")
    _paragraph(doc, 26, "A model replacement must be evaluated through the same frozen E2E matrix. Endpoint reachability and one successful dialogue establish operational availability, not semantic equivalence.")
    _paragraph(doc, 28, "Recommendations")
    _paragraph(doc, 29, "Keep vLLM as the preferred backend only after /v1/models and a named completion probe succeed.")
    _paragraph(doc, 30, "Use an explicit, launch-time fallback order and pin the selected model for the run. Record fallback reason and model identity in JSONL.")
    _paragraph(doc, 31, "Repeat shared failures and model-specific sensitive cases at least three times before changing stack code or prompt policy.")
    _paragraph(doc, 33, "The comparison supports a model-agnostic stack claim at the ownership and execution boundaries, with a qualification that model-mediated planning remains variable.")
    _paragraph(doc, 34, "The retained raw JSON, dataset, hypothesis registry, and HTML report make the attribution auditable for thesis Results and Discussion writing.")
    _paragraph(doc, 37, "Evidence semantics")
    _paragraph(doc, 38, "Deterministic tests validate implementation contracts. Live cases exercise endpoint reachability, ROS lifecycle, grounding, planning, execution, speech, and trace seams. The categories must remain separate in the thesis.")
    _paragraph(doc, 40, "NAO ROS4HRI runtime stack. Model invariance E2E comparison and raw JSON artifacts. Local project evidence pack, 21 July 2026.")
    _paragraph(doc, 41, "Robot runtime performance review. REPORT.md and REPORT.html for the frozen v34 image. Local project artifact, 21 July 2026.")
    _paragraph(doc, 42, "Model comparison dataset. dataset.csv, dataset.jsonl, and comparison_summary.json. Generated from the retained live questionnaires, 21 July 2026.")
    _paragraph(doc, 43, "Backend inventory. vLLM and Ollama endpoint probe with selected model identities. Local runtime evidence, 21 July 2026.")
    _paragraph(doc, 45, "Template note. The retained design-report visual system was used. The generated DOCX was rendered and checked after generation; raw evidence remains alongside this report.")

    doc.add_page_break()
    doc.add_heading("Comparative evidence tables", level=1)
    doc.add_paragraph("The following tables summarize the JSON outputs. Source-file paths and full case observations are retained in the accompanying dataset.")
    _table(
        doc,
        ["Case family", "Gemma4 31B", "Gemma4 cloud", "Nemotron 3 Super"],
        [
            ["Environment", "11/11", "11/11", "11/11"],
            ["Main", "17/20", "18/21", "17/19 + 1 degraded"],
            ["KB stress", "7/7", "7/7", "4/6"],
            ["Robustness", "3/5", "3/5", "3/5"],
            ["Fake deep all success", "9/9", "9/9", "8/9 + 1 degraded"],
            ["Capability extreme", "3/7", "3/7", "0/6 + 1 not scored"],
            ["Fail-once navigation", "diagnostic full run", "1/1", "1/1"],
        ],
        [2.1, 1.45, 1.45, 1.8],
    )

    doc.add_heading("Startup and endpoint conditions", level=1)
    _table(
        doc,
        ["Condition", "Observation", "Interpretation"],
        [
            ["vLLM inventory", "HTTP 000 at 10.7.138.215:8004", "Blocked semantic cell"],
            ["Ollama inventory", "11 advertised models", "Candidate source available"],
            ["Chatbot and planner preflight", "Pass for all three selected models", "Launch readiness established"],
            ["Required lifecycle nodes", "Active [3] in each scored cell", "Core graph ready"],
            ["NAOqi endpoint", "Connection timeout", "Isolated from alternative semantic cells"],
        ],
        [1.8, 3.2, 1.8],
    )

    doc.add_heading("Shared and model-specific outcomes", level=1)
    _table(
        doc,
        ["Evidence", "Finding", "Attribution posture"],
        [
            ["Gemma intersection", "Same generic pick, maximal delivery, missing-recipient, and extreme coverage failures", "Investigate stack or shared contract seams"],
            ["Nemotron variance", "KB wording omissions, partial report closure, planner coverage, and trace inconsistency", "Model/backend or observability candidates"],
            ["Fake recovery", "Targeted fail-once navigation passed for Nemotron and Gemma cloud", "Executor recovery seam remains viable"],
            ["Environment", "11/11 for all models", "No broad grounding outage supported"],
        ],
        [1.6, 3.45, 1.75],
    )

    doc.add_page_break()
    doc.add_heading("Seam hypothesis adjudication", level=1)
    _table(
        doc,
        ["Hypothesis", "Evidence", "Status"],
        [
            ["H-01 runtime wiring", "All Ollama preflights passed; vLLM unavailable", "Ollama route ruled out; vLLM blocked"],
            ["H-02 model policy", "Gemma intersection stable; Nemotron adds variance", "Supported"],
            ["H-03 grounding and KB", "Environment stable; Nemotron speech terms varied", "Broad outage rejected"],
            ["H-04 executor and recovery", "Fail-once navigation passed with replan", "Core supported; closure open"],
            ["H-05 observability", "Trace inconsistency and global timeouts retained", "Active limitation"],
            ["H-06 external service", "vLLM unavailable; Ollama preflight stable", "Backend condition confirmed"],
        ],
        [1.65, 3.7, 1.45],
    )
    doc.add_paragraph("The shared Gemma failure shape is the key adversarial result. It prevents a simple model-quality explanation and directs the next investigation toward planner admission, target normalization, clarification speech, and high-composition plan coverage. The Nemotron-only cases remain candidate model variance until the same prompts are repeated with stable trace and speech correlation.")

    doc.add_heading("Suggested TFM wording", level=1)
    doc.add_paragraph("The evaluation shows that the ROS4HRI stack preserves its principal ownership and execution contracts across the tested language-model backends. Generated target selection, multi-step plan coverage, clarification wording, and terminal report closure vary with the selected model. Shared failures across the two Gemma variants were retained as stack-seam hypotheses rather than attributed to model quality alone.")

    doc.add_heading("Reproducibility and handoff", level=1)
    _add_bullets(
        doc,
        [
            "Use analysis/comparison_summary.json for model-level tables and shared-failure intersections.",
            "Use analysis/dataset.csv for spreadsheet analysis and analysis/dataset.jsonl for programmatic TFM processing.",
            "Use raw per-model JSON when quoting exact speech, route, target, lineage, or timing evidence.",
            "Do not merge deterministic source-test counts into the live runtime stress denominator.",
            "When vLLM returns a stable model identity, repeat the same case matrix before promoting it as a replacement.",
        ],
    )

    doc.add_page_break()
    doc.add_heading("Appendix: provenance", level=1)
    doc.add_paragraph("The evidence pack preserves the endpoint probe, frozen image identity, startup snapshots, per-model questionnaire outputs, derived tables, methodology, and seam audit. The comparison is intentionally bounded to one frozen run per model. Repeated sensitive cells are the next acceptance gate for a model replacement.")
    _table(
        doc,
        ["Artifact", "Purpose"],
        [
            ["REPORT.md and REPORT.html", "Readable runtime review and thesis-facing discussion"],
            ["analysis/comparison_summary.json", "Model, suite, startup, and shared-failure summary"],
            ["analysis/dataset.csv", "Tabular per-case dataset with provenance"],
            ["analysis/dataset.jsonl", "Machine-readable per-case dataset"],
            ["raw/", "Authoritative JSON outputs for all model cells"],
            ["../model_invariance_seam_audit_2026-07-21.md", "Hypothesis registry and adversarial attribution"],
        ],
        [2.55, 4.25],
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    _replace_template_header(output)
    _replace_header_text(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--summary", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.template, args.summary, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
